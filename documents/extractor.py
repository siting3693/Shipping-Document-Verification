#!/usr/bin/env python3
"""Attachment text extraction supporting TXT, PDF, DOCX, and XLSX formats."""

import logging
from pathlib import Path
from typing import Any
import warnings

# Suppress pypdf verbose warnings on malformed stream bytes
warnings.filterwarnings("ignore", module="pypdf")
logging.getLogger("pypdf").setLevel(logging.ERROR)


def extract_attachment_text(att_path_str: str, max_excerpt_chars: int = 10000) -> dict[str, Any]:
    """Extract text from an attachment file.

    Returns:
        dict with keys:
            - status: "ok" or "unreadable"
            - error: None or error description
            - text: extracted text string (up to max_excerpt_chars)
            - page_count: number of pages / sheets
            - format: file extension (e.g. "pdf", "txt")
            - is_scanned: boolean indicating scanned/image-only without text
    """
    path = Path(att_path_str)
    suffix = path.suffix.lower().lstrip(".")

    if not path.exists():
        return {
            "status": "unreadable",
            "error": "file_not_found",
            "text": "",
            "page_count": 0,
            "format": suffix,
            "is_scanned": False,
        }

    try:
        size = path.stat().st_size
    except Exception as ex:
        return {
            "status": "unreadable",
            "error": f"stat_failed: {ex}",
            "text": "",
            "page_count": 0,
            "format": suffix,
            "is_scanned": False,
        }

    if size == 0:
        return {
            "status": "unreadable",
            "error": "empty_file",
            "text": "",
            "page_count": 0,
            "format": suffix,
            "is_scanned": False,
        }

    if suffix == "txt":
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
            return {
                "status": "ok",
                "error": None,
                "text": text[:max_excerpt_chars],
                "page_count": 1,
                "format": "txt",
                "is_scanned": False,
            }
        except Exception as ex:
            return {
                "status": "unreadable",
                "error": str(ex),
                "text": "",
                "page_count": 0,
                "format": "txt",
                "is_scanned": False,
            }

    elif suffix == "docx":
        try:
            import docx

            doc = docx.Document(str(path))
            parts: list[str] = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_txt:
                        parts.append(" | ".join(row_txt))
            text = "\n".join(parts)
            return {
                "status": "ok",
                "error": None,
                "text": text[:max_excerpt_chars],
                "page_count": 1,
                "format": "docx",
                "is_scanned": False,
            }
        except Exception as ex:
            return {
                "status": "unreadable",
                "error": f"docx_read_error: {ex}",
                "text": "",
                "page_count": 0,
                "format": "docx",
                "is_scanned": False,
            }

    elif suffix == "xlsx":
        try:
            import openpyxl

            wb = openpyxl.load_workbook(str(path), data_only=True)
            parts: list[str] = []
            for sheet in wb.sheetnames:
                parts.append(f"SHEET: {sheet}")
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        parts.append(" | ".join(row_vals))
            text = "\n".join(parts)
            return {
                "status": "ok",
                "error": None,
                "text": text[:max_excerpt_chars],
                "page_count": len(wb.sheetnames),
                "format": "xlsx",
                "is_scanned": False,
            }
        except Exception as ex:
            return {
                "status": "unreadable",
                "error": f"xlsx_read_error: {ex}",
                "text": "",
                "page_count": 0,
                "format": "xlsx",
                "is_scanned": False,
            }

    elif suffix == "pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(str(path))
            parts: list[str] = []
            page_count = len(reader.pages)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    parts.append(t)
            text = "\n".join(parts).strip()
            if not text:
                # PDF has pages, but no extractable text layer (scanned/image only)
                return {
                    "status": "unreadable",
                    "error": "scanned_or_image_only",
                    "text": "",
                    "page_count": page_count,
                    "format": "pdf",
                    "is_scanned": True,
                }
            return {
                "status": "ok",
                "error": None,
                "text": text[:max_excerpt_chars],
                "page_count": page_count,
                "format": "pdf",
                "is_scanned": False,
            }
        except Exception as ex:
            return {
                "status": "unreadable",
                "error": f"corrupt_or_garbled_pdf: {ex}",
                "text": "",
                "page_count": 0,
                "format": "pdf",
                "is_scanned": False,
            }

    return {
        "status": "unreadable",
        "error": f"unsupported_format: {suffix}",
        "text": "",
        "page_count": 0,
        "format": suffix,
        "is_scanned": False,
    }
