import io
import os
import re
from typing import Optional

def detect_document_type(text: str) -> str:
    """Detect document type from header text."""
    if not text:
        return 'UNREADABLE'
        
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if not lines:
        return 'UNREADABLE'
        
    first_few_lines = '\n'.join(lines[:10]).upper()
    
    # Check for wrong document types first
    if 'COMMERCIAL INVOICE' in first_few_lines:
        return 'COMMERCIAL_INVOICE'
    if 'PACKING LIST' in first_few_lines:
        return 'PACKING_LIST'
    if 'CERTIFICATE OF ORIGIN' in first_few_lines:
        return 'CERTIFICATE_OF_ORIGIN'
    
    # SI detection — "BILL OF LADING INSTRUCTION" and "BL INSTRUCTION" are actually SIs
    if 'SHIPPING INSTRUCTION' in first_few_lines:
        return 'SI'
    if 'BL INSTRUCTION' in first_few_lines and 'BILL OF LADING (DRAFT)' not in first_few_lines:
        return 'SI'
    if 'BILL OF LADING INSTRUCTION' in first_few_lines and 'BILL OF LADING (DRAFT)' not in first_few_lines:
        return 'SI'
    
    # BL detection
    if 'BILL OF LADING' in first_few_lines or 'B/L' in first_few_lines or 'WAYBILL' in first_few_lines:
        return 'BL'
    
    return 'UNKNOWN'

def _parse_text_content(text: str) -> dict:
    """Parse key:value pairs from text content.
    
    Handles formats:
    - 'Key: Value' (standard)
    - 'Key Value' (PDF format, for known keys)
    - Indented continuation lines
    - Separator lines (=====, -----)
    """
    raw_fields = {}
    lines = text.split('\n')
    current_key = None
    
    # Known field prefixes for colon-less matching (PDF format)
    known_keys = [
        'Shipper', 'SHIPPER', 'Consignee', 'CONSIGNEE', 'Notify', 'NOTIFY',
        'Notify Party', 'NOTIFY PARTY',
    ]
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        # Skip separator lines
        if stripped.startswith('===') or stripped.startswith('---'):
            continue
            
        # Check if line is indented (continuation of previous value)
        if (line.startswith(' ') or line.startswith('\t')) and not stripped.startswith(('Shipper', 'Consignee', 'Notify', 'Port', 'Load', 'Discharge', 'Container', 'Gross', 'TOTAL', 'No.')):
            if current_key:
                raw_fields[current_key] += " " + stripped
            continue
            
        # Check if it has a colon
        if ':' in line:
            parts = line.split(':', 1)
            key = parts[0].strip()
            value = parts[1].strip()
            raw_fields[key] = value
            current_key = key
        else:
            # Fallback for known prefixes without colons
            lower_line = line.strip().lower()
            prefixes = [
                'shipper/exporter', 'shipper (principal or seller)', 'shipper', 'exporter', 'seller',
                'consignee (non-negotiable)', 'consignee', 'to the order of', 'buyer',
                'notify party/intermediate consignee', 'notify party', 'notify',
                'port of loading (pol)', 'port of loading', 'load port', 'pol',
                'port of discharge (pod)', 'port of discharge', 'discharge port', 'pod',
                'no. of containers or packages', 'total containers', 'container count', 'no. of containers',
                'gross weight (kg)', 'gross weight', 'gross wt (kgs)', 'gross wt', 'total gross weight'
            ]
            matched = False
            for p in prefixes:
                if lower_line.startswith(p) and len(lower_line) > len(p) and lower_line[len(p)] == ' ':
                    key = line.strip()[:len(p)]
                    value = line.strip()[len(p):].strip()
                    raw_fields[key] = value
                    current_key = key
                    matched = True
                    break
            if not matched:
                current_key = None
            
    return raw_fields

def _parse_pdf(file_path: str, content_bytes: Optional[bytes] = None) -> dict:
    if content_bytes is None:
        try:
            with open(file_path, 'rb') as f:
                content_bytes = f.read()
        except Exception as e:
            return {
                'doc_type': 'UNREADABLE',
                'raw_fields': {},
                'raw_text': "",
                'parse_method': 'pdf',
                'error': f"Failed to read file: {str(e)}"
            }

    try:
        import pdfplumber
        import io
        text = ""
        with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        
        # If pdfplumber extracted nothing, it might be a weird PDF, try fallback
        if not text.strip():
            raise ValueError("No text extracted by pdfplumber")
            
        # Check for garbled OCR tokens (≥4 case transitions in a word).
        # pdfplumber sometimes garbles overlapping font paths; pypdfium2
        # reads the same PDFs cleanly.  This is parser-level quality
        # control — not a change to comparison logic.
        _has_garbled = False
        for token in text.split():
            alpha = ''.join(c for c in token if c.isalpha())
            if len(alpha) >= 5:
                transitions = sum(
                    1 for i in range(1, len(alpha))
                    if alpha[i - 1].isupper() != alpha[i].isupper()
                )
                if transitions >= 4:
                    _has_garbled = True
                    break
        if _has_garbled:
            raise ValueError("Garbled OCR tokens detected in pdfplumber output")
            
        doc_type = detect_document_type(text)
        raw_fields = _parse_text_content(text)
        return {
            'doc_type': doc_type,
            'raw_fields': raw_fields,
            'raw_text': text,
            'parse_method': 'pdf',
            'error': None
        }
    except Exception as e:
        # Fallback to pypdfium2 for corrupted PDFs like email_499_BL
        try:
            import pypdfium2 as pdfium
            text = ""
            pdf = pdfium.PdfDocument(content_bytes)
            for i in range(len(pdf)):
                page = pdf[i]
                text += page.get_textpage().get_text_range() + "\n"
            
            doc_type = detect_document_type(text)
            raw_fields = _parse_text_content(text)
            return {
                'doc_type': doc_type,
                'raw_fields': raw_fields,
                'raw_text': text,
                'parse_method': 'pdf_fallback',
                'error': None
            }
        except Exception as fallback_e:
            return {
                'doc_type': 'UNREADABLE',
                'raw_fields': {},
                'raw_text': "",
                'parse_method': 'pdf',
                'error': f"Primary error: {str(e)}, Fallback error: {str(fallback_e)}"
            }

def parse_document(file_path: str, content_bytes: Optional[bytes] = None) -> dict:
    """Parse a document file and return structured data.
    
    Args:
        file_path: Path to the file (used to determine type from extension)
        content_bytes: Optional raw bytes. If None, reads from file_path.
    
    Returns:
        {
            'doc_type': str,
            'raw_fields': dict[str, str],
            'raw_text': str,
            'parse_method': str,
            'error': str | None,
        }
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    result = {
        'doc_type': 'UNKNOWN',
        'raw_fields': {},
        'raw_text': '',
        'parse_method': '',
        'error': None
    }
    
    try:
        if ext == '.txt':
            result['parse_method'] = 'txt'
            if content_bytes:
                text = content_bytes.decode('utf-8', errors='replace')
            else:
                with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                    text = f.read()
            result['raw_text'] = text
            result['doc_type'] = detect_document_type(text)
            result['raw_fields'] = _parse_text_content(text)
            
        elif ext == '.pdf':
            return _parse_pdf(file_path, content_bytes)
            
        elif ext == '.docx':
            result['parse_method'] = 'docx'
            import docx
            
            if content_bytes:
                doc = docx.Document(io.BytesIO(content_bytes))
            else:
                doc = docx.Document(file_path)
                
            text_blocks = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_blocks.append(para.text)
                    
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        if len(row_data) == 2 and not row_data[0].endswith(':'):
                            text_blocks.append(f"{row_data[0]}: {row_data[1]}")
                        else:
                            text_blocks.append(" | ".join(row_data))
                        
            text = "\n".join(text_blocks)
            result['raw_text'] = text
            result['doc_type'] = detect_document_type(text)
            result['raw_fields'] = _parse_text_content(text)
            
        elif ext in ['.xlsx', '.xls']:
            result['parse_method'] = 'xlsx'
            import openpyxl
            
            if content_bytes:
                wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            else:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                
            text_blocks = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if row_data:
                        text_blocks.append(" ".join(row_data))
                        if len(row_data) == 2 and not row_data[0].endswith(':'):
                            text_blocks[-1] = f"{row_data[0]}: {row_data[1]}"
                                
            text = "\n".join(text_blocks)
            result['raw_text'] = text
            result['doc_type'] = detect_document_type(text)
            result['raw_fields'] = _parse_text_content(text)
            
        else:
            result['parse_method'] = 'unknown'
            result['doc_type'] = 'UNREADABLE'
            result['error'] = f'Unsupported file extension: {ext}'
            
    except Exception as e:
        result['doc_type'] = 'UNREADABLE'
        result['error'] = str(e)
        
    return result
